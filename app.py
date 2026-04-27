from flask import Flask, request, send_file, render_template
from PIL import Image, ImageOps
import os
from docx import Document
from docx.shared import Cm
import pillow_heif

pillow_heif.register_heif_opener()

app = Flask(__name__)

@app.route('/')
def index():
    return '''
    <h2>Upload Fotos</h2>
    <form method="post" action="/gerar" enctype="multipart/form-data">
        <input type="file" name="fotos" multiple>
        <button type="submit">Gerar Relatório</button>
    </form>
    '''

@app.route('/gerar', methods=['POST'])
def gerar():
    files = request.files.getlist('fotos')

    os.makedirs('temp', exist_ok=True)

    imagens = []
    nomes = []

    for i, f in enumerate(files, 1):
        path = f"temp/{i}.jpg"
        img = Image.open(f)
        img = ImageOps.exif_transpose(img)

        if img.mode != 'RGB':
            img = img.convert('RGB')

        w, h = img.size
        min_side = min(w, h)
        left = (w - min_side)//2
        top = (h - min_side)//2

        img = img.crop((left, top, left+min_side, top+min_side))
        img = img.resize((900,900))
        img.save(path)

        imagens.append(path)
        nomes.append(f.filename)

    doc = Document()
    foto = 1

    for i in range(0, len(imagens), 4):
        group = imagens[i:i+4]
        names = nomes[i:i+4]

        if i > 0:
            doc.add_page_break()

        table = doc.add_table(rows=2, cols=2)

        idx = 0
        for r in range(2):
            for c in range(2):
                cell = table.cell(r,c)

                if idx < len(group):
                    p = cell.paragraphs[0]
                    run = p.add_run()
                    run.add_picture(group[idx], width=Cm(7.5))

                    nome = names[idx].split('.')[0]

                    legenda = cell.add_paragraph()
                    legenda.add_run(f"Foto {foto:02d} - {nome}").bold = True

                    foto += 1
                    idx += 1

    doc.save("relatorio.docx")

    return send_file("relatorio.docx", as_attachment=True)

app.run(host='0.0.0.0', port=10000)
